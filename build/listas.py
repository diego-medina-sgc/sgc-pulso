# -*- coding: utf-8 -*-
"""Las listas de personas, leidas desde donde dice el panel de Fuentes.

    from listas import Listas
    L = Listas(cli)
    for lista in L.de_tipo("exalumnos"):
        filas = L.filas(lista)            # la hoja de la lista, fila por fila
        L.anotar(lista, len(filas), "ok")

POR QUE (14/9/2026)

Cada cargador tenia el id de su planilla escrito en el codigo, o leia un Excel
de la carpeta Descargas. El link que se veia en el panel "De donde salen los
nombres" no lo usaba nadie: Diego podia cambiarlo y la carga seguia leyendo la
planilla vieja. Ahora el link y la hoja viven en padron_fuentes
(sql/fuentes.sql) y todos los cargadores leen de aca.

La planilla se baja entera como .xlsx por la API de Drive -la de Sheets no
esta habilitada en el proyecto- y se lee con openpyxl, asi que la hoja se
elige por nombre y no por gid.
"""
import io
import os
import re
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
SA_PATH = os.path.join(HERE, "service-account.json")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def sheet_id(url):
    m = re.search(r"/d/([A-Za-z0-9_-]{20,})", url or "")
    return m.group(1) if m else None


class Listas:
    def __init__(self, cli):
        self.cli = cli
        self._tok = None
        self._libros = {}

    def token(self):
        if self._tok is None:
            from google.auth.transport.requests import Request
            from google.oauth2 import service_account
            cred = service_account.Credentials.from_service_account_file(
                SA_PATH, scopes=["https://www.googleapis.com/auth/drive.readonly"])
            cred.refresh(Request())
            self._tok = cred.token
        return self._tok

    def todas(self):
        return self.cli.select(
            "padron_fuentes",
            select="clave,titulo,url,hoja,persona_tipo,sede,historial")

    def de_tipo(self, tipo):
        """Las listas vigentes de ese tipo que tienen link."""
        return [l for l in self.todas()
                if l.get("persona_tipo") == tipo and not l.get("historial") and l.get("url")]

    def libro(self, lista):
        sid = sheet_id(lista.get("url"))
        if not sid:
            raise RuntimeError("la lista %s no tiene un link de planilla valido" % lista["clave"])
        if sid not in self._libros:
            import openpyxl
            req = urllib.request.Request(
                "https://www.googleapis.com/drive/v3/files/%s/export?mimeType=%s" % (sid, XLSX),
                headers={"Authorization": "Bearer " + self.token()})
            with urllib.request.urlopen(req, timeout=300) as r:
                self._libros[sid] = openpyxl.load_workbook(io.BytesIO(r.read()),
                                                           read_only=True, data_only=True)
        return self._libros[sid]

    def filas(self, lista, hoja=None):
        """Las filas no vacias de la hoja, como listas de celdas (None si vacia)."""
        wb = self.libro(lista)
        nombre = hoja or lista.get("hoja")
        if nombre not in wb.sheetnames:
            raise RuntimeError("la planilla de %s no tiene la hoja %r (tiene: %s)"
                               % (lista["clave"], nombre, ", ".join(wb.sheetnames)))
        return [list(r) for r in wb[nombre].iter_rows(values_only=True)
                if any(c not in (None, "") for c in r)]

    def anotar(self, lista, filas, resultado):
        try:
            self.cli.rpc("lista_leida", {"p_clave": lista["clave"], "p_filas": int(filas),
                                         "p_resultado": resultado})
        except Exception as e:
            print("  (no se pudo anotar la lectura de %s: %s)" % (lista["clave"], str(e)[:120]))


def texto(v):
    """Una celda como texto limpio: los numeros enteros sin '.0', las fechas ISO."""
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def indice(encabezado):
    """{nombre de columna normalizado: posicion}. Tolera saltos de linea y espacios."""
    out = {}
    for i, c in enumerate(encabezado):
        k = " ".join(str(c or "").replace("\n", " ").split()).lower()
        if k and k not in out:
            out[k] = i
    return out
